"""記事のアイキャッチ画像（1200×630 PNG）を PowerPoint から作る。

流れ:
  1. svg/*.svg（手書きの図版）を LibreOffice で PNG に起こす（PowerPoint 用の代替画像）
  2. python-pptx で article-images.pptx を組む（1スライド＝1記事）
  3. できた .pptx を開き直して、図版を **SVG のまま**埋め込み直す
     （PowerPoint 2016 以降の asvg:svgBlip 拡張。PNG は代替として残る）
  4. .pptx → PDF → 1200×630 の PNG に書き出して site/public/img/og/ に置く

  3 をやっているのは、PowerPoint で開いたときに図版がベクタのまま編集できるようにするため。
  python-pptx は SVG を直接 add_picture できないので、保存後の OOXML を触っている。

使い方:
    python tools/article-images/build.py

前提: LibreOffice（soffice.exe）、python-pptx、PyMuPDF。
"""

import json
import os
import re
import shutil
import subprocess
import sys
import zipfile

import fitz  # PyMuPDF
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
SVG_DIR = os.path.join(HERE, "svg")
WORK = os.path.join(HERE, "_work")
PPTX = os.path.join(HERE, "article-images.pptx")
OUT_DIR = os.path.join(REPO, "site", "public", "img", "og")

SOFFICE = (
    r"C:\Program Files\LibreOffice\program\soffice.exe"
    if os.path.exists(r"C:\Program Files\LibreOffice\program\soffice.exe")
    else "soffice"
)
# 実行中の LibreOffice と衝突しないよう、専用のユーザープロファイルを使う
LO_PROFILE = "file:///" + os.path.join(WORK, "loprofile").replace("\\", "/")

# 色は site/public/styles.css のカスタムプロパティに合わせてある。
# 画面（紺のヘッダ・表・カテゴリラベル）と画像の色が食い違わないようにするため。
NAVY = RGBColor(0x01, 0x41, 0x72)   # --navy
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AZURE = RGBColor(0x8E, 0xC5, 0xE8)  # --link (#0077c6) を紺の上で読める明るさにしたもの
WARN = RGBColor(0xE7, 0xCD, 0x7A)   # --warn-line。図の黄色い破線と、それを指すラベルに使う
JP_FONT = "Noto Sans CJK JP" if os.name != "nt" else "Yu Gothic UI"

# 1200×630 px を 96dpi として inch に直した値
SLIDE_W = Inches(12.5)
SLIDE_H = Inches(6.5625)

# 図版（SVG）を置く枠。ラベルの座標変換もこの枠を使うので、動かすときは片方だけ直さない。
# ⚠️ SVG の縦横比とこの枠の縦横比は一致していない（縦横で別々に伸びる）。
#    だからラベルの座標変換も x と y で別々の倍率をかける。
FIG_L, FIG_T = Inches(6.62), Inches(0.62)
FIG_W, FIG_H = Inches(5.21), Inches(4.69)

SLIDES = [
    {
        "svg": "fuyou-no-kabe.svg",
        "png": "fuyou-no-kabe.png",
        "category": "税と社会保険",
        "title": "扶養の壁",
        "title_size": 46,
        "subtitle": "103・106・130・150・160 は\n何の壁なのか",
        # 壁の下に「何の壁か」を1本ずつ置く。色分け（白＝税／青＝社会保険）は
        # ラベルで言い換えられるので凡例は置かない。
        "labels": [
            {"at": (70, 428),  "text": "所得税\n〜令和6年分", "size": 10, "align": "center", "width": Inches(0.94)},
            {"at": (160, 428), "text": "社会保険\n適用拡大", "size": 10, "align": "center", "width": Inches(0.94), "color": AZURE},
            {"at": (250, 428), "text": "健康保険の\n被扶養者", "size": 10, "align": "center", "width": Inches(0.94), "color": AZURE},
            {"at": (340, 428), "text": "配偶者\n特別控除", "size": 10, "align": "center", "width": Inches(0.94)},
            {"at": (430, 428), "text": "所得税\n令和7年分〜", "size": 10, "align": "center", "width": Inches(0.94)},
        ],
        "note": "単位：万円　白＝税の壁　青＝社会保険の壁",
    },
    {
        "svg": "furusato-nozei-jogen.svg",
        "png": "furusato-nozei-jogen.png",
        "category": "税と社会保険",
        "title": "ふるさと納税の上限額",
        "title_size": 34,
        "subtitle": "決めているのは\n総務省の3本の計算式",
        # ラベルで名前が付いたので凡例は置かない（色と名前を二度説明することになる）
        "labels": [
            {"at": (212, 345), "text": "(1) 所得税から", "size": 15},
            {"at": (212, 263), "text": "(2) 住民税の基本分", "size": 15},
            {"at": (212, 169), "text": "(3) 住民税の特例分", "size": 15},
            {"at": (212, 74), "text": "住民税所得割額の20％\n＝ここが天井",
             "size": 13, "color": WARN},
        ],
        "note": "破線の上にはみ出した分は控除されず、自己負担として残る",
    },
    {
        "svg": "koukou-mushouka.svg",
        "png": "koukou-mushouka.png",
        "category": "教育費",
        "title": "「高校無償化」の中身",
        "title_size": 34,
        "subtitle": "制度名は高等学校等就学支援金\n授業料以外は別の給付金",
        # 「何が2つに割れているのか」と「矢印の行き先が何か」を図の中で言う。
        "labels": [
            {"at": (250, 69),  "text": "高校でかかる費用", "size": 13, "align": "center", "width": Inches(2.2)},
            {"at": (132, 103), "text": "授業料", "size": 15, "align": "center", "width": Inches(1.7)},
            {"at": (368, 103), "text": "授業料以外", "size": 15, "align": "center", "width": Inches(1.7),
             "color": AZURE},
            {"at": (132, 388), "text": "高等学校等\n就学支援金", "size": 11, "align": "center", "width": Inches(1.7)},
            {"at": (368, 388), "text": "高校生等\n奨学給付金", "size": 11, "align": "center", "width": Inches(1.7),
             "color": AZURE},
        ],
        "note": "矢印の先＝それを受け止める制度\n数字は私立高校の年額（円）",
    },
    {
        "svg": "daigaku-mushouka.svg",
        "png": "daigaku-mushouka.png",
        "category": "教育費",
        "title": "「大学無償化」の中身",
        "title_size": 34,
        "subtitle": "高等教育の修学支援新制度\n無償化されたのは減免だけ",
        # 2本の支援に名前を付け、「所得制限が外れたのは減免だけ」を図の中で言う。
        # 箱の中のラベルは背景が明るいので NAVY。
        "labels": [
            {"at": (250, 63),  "text": "高等教育の修学支援新制度", "size": 13, "align": "center",
             "width": Inches(2.8)},
            {"at": (137, 185), "text": "授業料等減免", "size": 14, "align": "center",
             "width": Inches(1.7), "color": NAVY},
            {"at": (363, 150), "text": "給付型奨学金", "size": 14, "align": "center",
             "width": Inches(1.7), "color": NAVY},
            {"at": (137, 422), "text": "多子世帯は\n所得制限なし", "size": 10, "align": "center",
             "width": Inches(1.7)},
            {"at": (363, 422), "text": "収入の基準が\n残っている", "size": 10, "align": "center",
             "width": Inches(1.7), "color": AZURE},
        ],
        "note": "数字は私立大学の年額（給付型は自宅外）",
    },
    {
        "svg": "kougaku-ryouyouhi.svg",
        "png": "kougaku-ryouyouhi.png",
        "category": "税と社会保険",
        "title": "高額療養費の年間上限",
        "title_size": 34,
        "subtitle": "2026年8月から\n年単位の上限額が入った",
        # 軸と破線を注記で説明していたが、図の中で直に言えるようになったので移した。
        "labels": [
            {"at": (44, 122),  "text": "年間上限", "size": 12, "color": WARN},
            # ⚠️ 階段の中に収める。左に寄せると白の外へ出て NAVY の字が読めなくなる。
            {"at": (250, 355), "text": "実際に払う自己負担", "size": 13, "align": "center",
             "width": Inches(2.0), "color": NAVY},
            {"at": (372, 62),  "text": "払わずに済む分", "size": 11, "align": "center",
             "width": Inches(1.6)},
            {"at": (60, 430),  "text": "8月", "size": 10, "align": "center", "width": Inches(0.8)},
            {"at": (432, 430), "text": "翌年7月", "size": 10, "align": "center", "width": Inches(0.9)},
        ],
        "note": "1段＝1か月ぶんの自己負担。破線から先は払わなくてよい",
    },
    {
        "svg": "jidouteate.svg",
        "png": "jidouteate.png",
        "category": "教育費",
        "title": "児童手当の「第3子」",
        "title_size": 34,
        "subtitle": "3人いても\n第3子とは限らない",
        # 何番目の子かを箱の中で直に言う。3人目だけ背景が白なので NAVY。
        "labels": [
            {"at": (300, 98),  "text": "1人目", "size": 14, "align": "center",
             "width": Inches(1.8), "color": AZURE},
            {"at": (300, 218), "text": "2人目", "size": 14, "align": "center",
             "width": Inches(1.8), "color": AZURE},
            {"at": (300, 338), "text": "3人目＝加算の対象", "size": 13, "align": "center",
             "width": Inches(1.9), "color": NAVY},
        ],
        "note": "点線＝経済的負担と確認書があれば数に入る上の子\n左の点＝数える順番（上の子から数える）",
    },
    {
        "svg": "iryouhi-koujo.svg",
        "png": "iryouhi-koujo.png",
        "category": "税と社会保険",
        "title": "医療費控除の下限額",
        "title_size": 34,
        "subtitle": "10万円とは限らない\n引く順番も決まっている",
        # 3段それぞれに名前を付け、破線が何かも図の中で言う。
        "labels": [
            {"at": (312, 150), "text": "控除の対象\nになる分", "size": 12, "width": Inches(1.8)},
            # ⚠️ 破線の右端（x=302）より右に置く。重ねると字が線に食われる。
            {"at": (318, 248), "text": "ここまでが引かれる", "size": 10, "width": Inches(1.8),
             "color": WARN},
            {"at": (312, 285), "text": "適用下限額", "size": 11, "width": Inches(1.8), "color": AZURE},
            {"at": (312, 362), "text": "保険金などで\n補てんされた分", "size": 10, "width": Inches(1.8),
             "color": AZURE},
        ],
        "note": "縦＝1年間に支払った医療費の積み上がり",
    },
    {
        "svg": "jutaku-loan-koujo.svg",
        "png": "jutaku-loan-koujo.png",
        "category": "税と社会保険",
        "title": "住宅ローン減税の区分",
        "title_size": 34,
        "subtitle": "省エネ基準を満たさない\n新築は支援対象外",
        # 1段が2本（新築＝白／既存＝青）なので凡例は残す。区分の名前は図の中に置く。
        "legend": [("新築住宅", WHITE), ("既存住宅（中古）", AZURE)],
        "labels": [
            {"at": (340, 93),  "text": "長期優良住宅\n低炭素住宅", "size": 11, "width": Inches(1.8)},
            {"at": (340, 189), "text": "ZEH水準\n省エネ住宅", "size": 11, "width": Inches(1.8)},
            {"at": (340, 285), "text": "省エネ基準\n適合住宅", "size": 11, "width": Inches(1.8)},
            {"at": (340, 381), "text": "その他住宅\n新築は支援対象外", "size": 11, "width": Inches(1.8),
             "color": WARN},
        ],
        "note": "横の長さ＝借入限度額（令和8年〜令和12年に入居した場合）",
    },
    {
        "svg": "shogakukin-henkan.svg",
        "png": "shogakukin-henkan.png",
        "category": "教育費",
        "title": "奨学金の返還が苦しいとき",
        "title_size": 30,
        "subtitle": "減額返還も猶予も\n返す総額は減らない",
        # 段の順番を注記で書いていたが、図の中で直に名乗れるようになったので移した。
        "legend": [("返す分", WHITE), ("待つ期間", AZURE)],
        "labels": [
            {"at": (358, 77),  "text": "元の返還", "size": 12, "width": Inches(1.7)},
            {"at": (358, 167), "text": "減額返還", "size": 12, "width": Inches(1.7)},
            {"at": (358, 257), "text": "返還期限猶予", "size": 12, "width": Inches(1.7)},
            {"at": (358, 347), "text": "返還免除", "size": 12, "width": Inches(1.7), "color": WARN},
        ],
        "note": "白の合計＝返す総額。減るのは免除だけ",
    },
    {
        "svg": "hoiku-mushouka.svg",
        "png": "hoiku-mushouka.png",
        "category": "教育費",
        "title": "「保育の無償化」の中身",
        "title_size": 34,
        "subtitle": "無料になるのは利用料だけ\n0〜2歳児クラスは原則対象外",
        # 列名（何の費目か）と行名（どの年齢か）を図の中で名乗らせる。
        # ⚠️ 棒グラフではないので、区画の幅で量を語らせない（SVG のコメント参照）。
        "labels": [
            {"at": (168, 42), "text": "利用料", "size": 11, "align": "center",
             "width": Inches(0.95), "color": WARN},
            {"at": (268, 42), "text": "食材料費", "size": 11, "align": "center",
             "width": Inches(0.95)},
            {"at": (368, 42), "text": "通園送迎費", "size": 11, "align": "center",
             "width": Inches(0.95)},
            {"at": (468, 42), "text": "行事費", "size": 11, "align": "center",
             "width": Inches(0.95)},
            {"at": (108, 122), "text": "3〜5歳児\nクラス", "size": 11, "align": "right",
             "width": Inches(1.2)},
            {"at": (108, 270), "text": "0〜2歳児\nクラス", "size": 11, "align": "right",
             "width": Inches(1.2)},
        ],
        "note": "白＝払う分　黄色い破線＝無償化で払わなくなる分\n0〜2歳児クラスは住民税非課税世帯だけ利用料が無償",
    },
    {
        "svg": "seimeihoken-hikazei.svg",
        "png": "seimeihoken-hikazei.png",
        "category": "相続",
        "title": "生命保険金の非課税枠",
        "title_size": 34,
        "subtitle": "受取人が相続人でないと\n1円も使えない",
        # 受取人の名前と、矢印の行き先が何かを図の中で言う。
        "labels": [
            {"at": (66, 76),  "text": "配偶者", "size": 13, "align": "center",
             "width": Inches(0.92), "color": NAVY},
            {"at": (66, 146), "text": "子", "size": 13, "align": "center",
             "width": Inches(0.92), "color": NAVY},
            {"at": (66, 264), "text": "孫", "size": 13, "align": "center",
             "width": Inches(0.92), "color": NAVY},
            {"at": (378, 111), "text": "非課税枠に入る", "size": 14, "align": "center",
             "width": Inches(2.1), "color": NAVY},
            {"at": (378, 264), "text": "枠の外\n非課税の適用なし", "size": 11, "align": "center",
             "width": Inches(2.1), "color": WARN},
        ],
        "note": "枠の大きさ＝500万円×法定相続人の数\n箱の大小は金額を表していない",
    },
    {
        "svg": "izoku-nenkin.svg",
        "png": "izoku-nenkin.png",
        "category": "年金",
        "title": "遺族基礎年金が止まる日",
        "title_size": 32,
        "subtitle": "子がいない世帯には\nはじめから出ない",
        # 行が何の世帯かと、帯が何かを図の中で言う。
        "labels": [
            {"at": (52, 62),  "text": "子がいる世帯", "size": 12},
            {"at": (52, 170), "text": "子がいない世帯", "size": 12},
            {"at": (184, 115), "text": "遺族基礎年金が出る", "size": 12, "align": "center",
             "width": Inches(2.0), "color": NAVY},
            {"at": (412, 115), "text": "止まる", "size": 11, "align": "center",
             "width": Inches(1.3), "color": WARN},
            {"at": (176, 223), "text": "はじめから出ない", "size": 11, "align": "center",
             "width": Inches(2.0), "color": WARN},
            {"at": (320, 318), "text": "子が18歳到達年度の末日", "size": 10, "align": "center",
             "width": Inches(2.4), "color": WARN},
        ],
        "note": "横＝時間　白＝出る期間　黄色い破線＝出ない期間\n帯の長さは期間の目安で、金額は表していない",
    },
    {
        "svg": "ideco-jougen.svg",
        "png": "ideco-jougen.png",
        "category": "資産形成",
        "title": "iDeCoの掛金上限",
        "title_size": 34,
        "subtitle": "2026年12月1日から\n会社員の差が無くなる",
        # 棒の名前は左の組にだけ置く（右も同じ並び）。右は「揃う」ことだけ言う。
        "labels": [
            {"at": (68, 114),  "text": "第1号", "size": 10, "align": "center",
             "width": Inches(0.78)},
            {"at": (148, 226), "text": "第2号\n企業年金なし", "size": 9, "align": "center",
             "width": Inches(0.78)},
            {"at": (228, 233), "text": "第2号\n企業年金あり", "size": 9, "align": "center",
             "width": Inches(0.78)},
            {"at": (398, 92),  "text": "第2号は同じ額に", "size": 11, "align": "center",
             "width": Inches(1.8), "color": WARN},
            {"at": (148, 322), "text": "2026年11月30日まで", "size": 10, "align": "center",
             "width": Inches(2.0)},
            {"at": (398, 322), "text": "2026年12月1日から", "size": 10, "align": "center",
             "width": Inches(2.0)},
        ],
        "note": "縦＝1か月の掛金上限（円）\n第1号＝自営業者等　第2号＝会社員・公務員",
    },
    {
        "svg": "kuriage-kurisage.svg",
        "png": "kuriage-kurisage.png",
        "category": "年金",
        "title": "年金の繰上げ・繰下げ",
        "title_size": 34,
        "subtitle": "繰上げは取り消せず\n繰下げでも増えない部分がある",
        # 年齢は棒の下、増減は棒の中、破線の意味は上に置く。
        "labels": [
            {"at": (98, 322),  "text": "60歳", "size": 12, "align": "center", "width": Inches(0.9)},
            {"at": (258, 322), "text": "65歳", "size": 12, "align": "center", "width": Inches(0.9)},
            {"at": (418, 322), "text": "75歳", "size": 12, "align": "center", "width": Inches(0.9)},
            # ⚠️ 棒の幅は 76px（約0.76インチ）しかない。1行で「84％増える」を置くと
            #    箱からはみ出し、紺地に紺の字が乗って読めなくなる（実際にそうなった）。
            {"at": (98, 254),  "text": "24％\n減る", "size": 10, "align": "center",
             "width": Inches(0.72), "color": NAVY},
            {"at": (258, 245), "text": "基準", "size": 11, "align": "center",
             "width": Inches(0.72), "color": NAVY},
            {"at": (418, 196), "text": "84％\n増える", "size": 10, "align": "center",
             "width": Inches(0.72), "color": NAVY},
            {"at": (418, 44),  "text": "加給年金額・振替加算額\nここは増えない", "size": 10,
             "align": "center", "width": Inches(2.2), "color": WARN},
        ],
        "note": "縦＝年金額（65歳を100とした割合）\n黄色い破線＝繰上げでも繰下げでも増減しない部分",
    },
    {
        "svg": "taishoku-shotoku.svg",
        "png": "taishoku-shotoku.png",
        "category": "資産形成",
        "title": "退職所得控除の折れ点",
        "title_size": 34,
        "subtitle": "勤続20年を境に\n1年あたり40万円→70万円",
        # 2つの傾きに名前を付け、折れ点が何年かを軸の下で言う。
        "labels": [
            # ⚠️ 折れ線の下（y=272）に置くと線が字を横切った。傾きがゆるい区間なので
            #    上の余白へ逃がす。
            {"at": (162, 232), "text": "1年あたり40万円", "size": 11, "align": "center",
             "width": Inches(1.7)},
            {"at": (372, 104), "text": "1年あたり70万円", "size": 11, "align": "center",
             "width": Inches(1.7)},
            {"at": (280, 324), "text": "勤続20年", "size": 11, "align": "center",
             "width": Inches(1.3), "color": WARN},
        ],
        "note": "横＝勤続年数　縦＝退職所得控除額\n目盛りの数字は置いていない（額は本文の表）",
    },
    {
        "svg": "shoukibo-takuchi.svg",
        "png": "shoukibo-takuchi.png",
        "category": "相続",
        "title": "小規模宅地等の特例",
        "title_size": 34,
        "subtitle": "誰が相続するかで\n満たす要件の数が変わる",
        # 取得者の名前を箱の中に、要件の中身を四角のとなりに置く。
        "labels": [
            {"at": (76, 77),  "text": "配偶者", "size": 12, "align": "center",
             "width": Inches(1.12), "color": NAVY},
            {"at": (76, 167), "text": "同居していた\n親族", "size": 10, "align": "center",
             "width": Inches(1.12), "color": NAVY},
            {"at": (76, 257), "text": "それ以外の\n親族", "size": 10, "align": "center",
             "width": Inches(1.12), "color": NAVY},
            {"at": (190, 77),  "text": "要件なし", "size": 13, "width": Inches(1.6)},
            {"at": (268, 167), "text": "申告期限まで\n居住し、所有する", "size": 10,
             "width": Inches(1.9)},
            {"at": (293, 220), "text": "(1)から(6)をすべて満たす", "size": 10, "align": "center",
             "width": Inches(2.4), "color": WARN},
        ],
        "note": "四角の数＝満たすべき要件の数\n面積や金額は表していない",
    },
    {
        "svg": "furusato-onestop.svg",
        "png": "furusato-onestop.png",
        "category": "税と社会保険",
        "title": "ふるさと納税ワンストップ特例",
        "title_size": 30,
        "subtitle": "控除は全額・翌年度の住民税\n確定申告すると申請は無効",
        "labels": [
            {"at": (73, 62),  "text": "確定申告", "size": 12, "align": "center",
             "width": Inches(1.1)},
            {"at": (183, 62), "text": "ワンストップ", "size": 12, "align": "center",
             "width": Inches(1.3)},
            {"at": (73, 308), "text": "所得税", "size": 12, "align": "center",
             "width": Inches(1.0), "color": NAVY},
            {"at": (73, 160), "text": "住民税", "size": 12, "align": "center",
             "width": Inches(1.0), "color": NAVY},
            {"at": (183, 200), "text": "全額\n住民税", "size": 13, "align": "center",
             "width": Inches(1.0), "color": NAVY},
            {"at": (250, 48), "text": "確定申告で無効", "size": 12, "color": WARN},
            {"at": (72, 430), "text": "5団体以内", "size": 11, "align": "center",
             "width": Inches(1.4)},
            {"at": (153, 430), "text": "6団体〜", "size": 11, "align": "center",
             "width": Inches(0.9), "color": WARN},
        ],
        "note": "白＝所得税　青＝住民税　黄＝無効・対象外の目印",
    },
    {
        "svg": "kafu-nenkin.svg",
        "png": "kafu-nenkin.png",
        "category": "年金",
        "title": "寡婦年金・死亡一時金",
        "title_size": 34,
        "subtitle": "遺族基礎年金が出ない世帯向けの\n第1号独自給付",
        "labels": [
            {"at": (40, 28),  "text": "遺族基礎年金が出ない", "size": 12,
             "width": Inches(2.6), "color": WARN},
            {"at": (40, 188), "text": "寡婦年金", "size": 13},
            {"at": (52, 300), "text": "60歳", "size": 11, "align": "center",
             "width": Inches(0.7)},
            {"at": (180, 300), "text": "65歳", "size": 11, "align": "center",
             "width": Inches(0.7)},
            {"at": (280, 172), "text": "死亡一時金", "size": 12,
             "width": Inches(1.4)},
            {"at": (40, 340), "text": "第1号被保険者だけの給付\n（厚生年金には無い）", "size": 11,
             "width": Inches(3.0), "color": AZURE},
        ],
        "note": "黄の空枠＝遺族基礎が出ない　白帯＝寡婦年金（60〜65歳）\n青箱＝死亡一時金（一回）　OR＝どちらか一方",
    },
    {
        "svg": "gakusei-nofu-tokurei.svg",
        "png": "gakusei-nofu-tokurei.png",
        "category": "年金",
        "title": "学生納付特例",
        "title_size": 36,
        "subtitle": "受給資格期間には入る\n年金額には入らない",
        "labels": [
            {"at": (150, 12), "text": "受給資格期間", "size": 12, "align": "center",
             "width": Inches(1.35)},
            {"at": (310, 12), "text": "年金額への反映", "size": 12, "align": "center",
             "width": Inches(1.35)},
            {"at": (28, 110), "text": "納付", "size": 14, "align": "center",
             "width": Inches(1.0)},
            {"at": (28, 212), "text": "学生納付\n特例", "size": 12, "align": "center",
             "width": Inches(1.0), "color": WARN},
            {"at": (28, 330), "text": "未納", "size": 14, "align": "center",
             "width": Inches(1.0)},
        ],
        "note": "白の実線＝入る　黄色い破線の空枠＝入らない\n○×はリーフレットの表と同じ構図",
    },
    {
        # 記事に eyecatch が無いときの既定の og:image（site.json の defaultOgImage）
        "kind": "site",
        "svg": "brand-mark.svg",
        "png": "site.png",
        "title": "家計の制度ログ",
        "subtitle": "お金の制度を、公式の原文で確かめて書く",
    },
]


def run_soffice(args):
    cmd = [SOFFICE, "--headless", "--norestore", f"-env:UserInstallation={LO_PROFILE}"] + args
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=300)
    if r.returncode != 0:
        raise RuntimeError(f"soffice failed: {r.returncode}\n{r.stdout}\n{r.stderr}")
    return r.stdout


def svg_to_png(svg_path, out_dir, width, height):
    """LibreOffice で SVG を透過 PNG に起こす。"""
    opts = {
        "PixelWidth": {"type": "long", "value": width},
        "PixelHeight": {"type": "long", "value": height},
        "Translucent": {"type": "boolean", "value": True},
    }
    run_soffice([
        "--convert-to",
        "png:draw_png_Export:" + json.dumps(opts, separators=(",", ":")),
        "--outdir",
        out_dir,
        svg_path,
    ])
    png = os.path.join(out_dir, os.path.splitext(os.path.basename(svg_path))[0] + ".png")
    if not os.path.exists(png):
        raise RuntimeError(f"PNG が作られませんでした: {png}")
    return png


def add_text(slide, left, top, width, height, runs, size, bold=False,
             color=WHITE, spacing=1.25, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, line in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = JP_FONT
    return box


def svg_viewbox(svg_name):
    """SVG の viewBox（幅・高さ）を返す。ラベルの座標変換に使う。"""
    head = open(os.path.join(SVG_DIR, svg_name), encoding="utf-8").read(2000)
    m = re.search(r'viewBox="\s*[\d.+-]+\s+[\d.+-]+\s+([\d.]+)\s+([\d.]+)', head)
    if not m:
        raise RuntimeError(f"{svg_name}: viewBox が読めません")
    return float(m.group(1)), float(m.group(2))


# ラベルを置く箱の幅。長いラベルはこの幅で折り返す。
LABEL_BOX_W = Inches(1.9)


def add_figure_labels(slide, spec):
    """図の要素のとなりに日本語のラベルを置く。

    **なぜ必要か**: SVG には日本語を入れない決まりなので（ラスタライズをフォントに
    依存させないため）、図の中の要素に名前が付けられなかった。その結果、読者は
    「図形 → 右下の凡例 → さらに小さい注記」と3往復しないと図を読めなかった。
    PowerPoint のテキストボックスなら日本語を置けるので、ここで図の上に重ねる。

    位置は **SVG の viewBox 座標**で書く。図を描いた座標系のまま指定できるので、
    SVG を直したときにラベルもそのまま追従させやすい。

    spec["labels"] の各要素:
        {"at": (x, y), "text": "授業料",
         "align": "left" | "right" | "center",   # at がテキストのどちら側か（既定 left）
         "size": 14, "color": WHITE, "bold": True}
    """
    labels = spec.get("labels")
    if not labels:
        return
    vb_w, vb_h = svg_viewbox(spec["svg"])
    for lb in labels:
        x, y = lb["at"]
        # ⚠️ x と y で別々の倍率をかける（図版は縦横比を保たずに枠へ引き伸ばされている）
        cx = FIG_L + int(FIG_W * (x / vb_w))
        cy = FIG_T + int(FIG_H * (y / vb_h))
        size = lb.get("size", 14)
        align = lb.get("align", "left")
        w = lb.get("width", LABEL_BOX_W)
        if align == "right":
            left, pp = cx - w, PP_ALIGN.RIGHT
        elif align == "center":
            left, pp = cx - int(w / 2), PP_ALIGN.CENTER
        else:
            left, pp = cx, PP_ALIGN.LEFT
        # at を文字の縦中央に合わせる（テキストボックスの原点は左上なので半行ぶん上げる）
        top = cy - Pt(size * 0.72)
        add_text(slide, left, top, w, Pt(size * 1.6),
                 lb["text"].split(chr(10)), size,
                 bold=lb.get("bold", True), color=lb.get("color", WHITE),
                 spacing=1.15, align=pp)


def build_pptx(png_paths):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    for spec in SLIDES:
        slide = prs.slides.add_slide(blank)

        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = NAVY

        # 左上のブランドの3本線（favicon と同じ「公式ページの本文」の見立て）
        for i, w in enumerate((2.6, 2.1, 1.6)):
            bar = slide.shapes.add_shape(
                5,  # MSO_SHAPE.ROUNDED_RECTANGLE
                Inches(0.78), Inches(0.62 + i * 0.235), Inches(w), Inches(0.075),
            )
            bar.fill.solid()
            bar.fill.fore_color.rgb = WHITE
            bar.fill.transparency = 0.72
            bar.line.fill.background()
            bar.shadow.inherit = False

        if spec.get("kind") == "site":
            add_text(slide, Inches(0.78), Inches(2.30), Inches(5.75), Inches(1.1),
                     [spec["title"]], 54, bold=True)
            add_text(slide, Inches(0.78), Inches(3.55), Inches(6.2), Inches(0.6),
                     [spec["subtitle"]], 21, color=WHITE)
            add_text(slide, Inches(0.78), Inches(5.62), Inches(5.6), Inches(0.4),
                     ["kakei.nexeed-lab.com"], 19, color=AZURE)
        else:
            add_text(slide, Inches(0.78), Inches(1.62), Inches(5.4), Inches(0.45),
                     [spec["category"]], 18, bold=True, color=AZURE)
            add_text(slide, Inches(0.78), Inches(2.12), Inches(5.75), Inches(1.0),
                     [spec["title"]], spec["title_size"], bold=True)
            add_text(slide, Inches(0.78), Inches(3.28), Inches(5.6), Inches(1.4),
                     spec["subtitle"].split(chr(10)), 26, spacing=1.35, color=WHITE)
            add_text(slide, Inches(0.78), Inches(5.62), Inches(5.6), Inches(0.4),
                     ["家計の制度ログ　kakei.nexeed-lab.com"], 17, color=AZURE)

        pic = slide.shapes.add_picture(
            png_paths[spec["svg"]], FIG_L, FIG_T, FIG_W, FIG_H,
        )
        spec["_pic_name"] = pic.name

        # 図の中のラベル。SVG の viewBox 座標で位置を指定する（add_figure_labels を参照）
        add_figure_labels(slide, spec)

        # 図版の凡例（日本語なのでフォントの都合上 SVG ではなくスライド側に置く）
        x = Inches(6.75)
        for label, color in spec.get("legend", []):
            chip = slide.shapes.add_shape(5, x, Inches(5.46), Inches(0.17), Inches(0.17))
            chip.fill.solid()
            chip.fill.fore_color.rgb = color
            chip.line.fill.background()
            chip.shadow.inherit = False
            add_text(slide, x + Inches(0.28), Inches(5.42), Inches(2.3), Inches(0.3),
                     [label], 15, color=WHITE)
            x += Inches(2.45)

        # note は改行で複数行にできる。図の軸が何かを書くために2行必要になった。
        if spec.get("note"):
            add_text(slide, Inches(6.62), Inches(5.88), Inches(5.21), Inches(0.3),
                     spec["note"].split(chr(10)), 13, color=AZURE)

    prs.save(PPTX)
    return PPTX


SVG_NS = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
SVG_EXT_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"


def embed_svg(pptx_path):
    """保存済み .pptx を開き直して、各スライドの図版に SVG 本体を埋め込む。

    PowerPoint は <a:blip> の拡張 asvg:svgBlip に SVG を持たせ、r:embed の PNG を
    代替として残す。python-pptx は SVG を扱えないので、ここは OOXML を直接触る。
    """
    tmp = pptx_path + ".tmp"
    with zipfile.ZipFile(pptx_path) as zin:
        names = zin.namelist()
        parts = {n: zin.read(n) for n in names}

    # [Content_Types].xml に svg の既定を足す
    ct = parts["[Content_Types].xml"].decode("utf-8")
    if 'Extension="svg"' not in ct:
        ct = ct.replace("<Default", '<Default Extension="svg" ContentType="image/svg+xml"/><Default', 1)
        parts["[Content_Types].xml"] = ct.encode("utf-8")

    # ⚠️ 辞書順で並べてはいけない。slide1, slide10, slide2 ... の順になり、
    #    スライドが10枚を超えた瞬間に SVG が1枚ずつずれる（2026-09-04 に実際に起きた）。
    #    ずれても例外は出ず、書き出した PNG に別の記事の図が入るだけなので、目で見るまで気づけない。
    slide_names = sorted(
        (n for n in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
        key=lambda n: int(re.search(r"slide(\d+)\.xml", n).group(1)),
    )
    if len(slide_names) != len(SLIDES):
        raise RuntimeError(f"スライド数が合いません: {len(slide_names)} != {len(SLIDES)}")

    for slide_name, spec in zip(slide_names, SLIDES):
        svg_src = os.path.join(SVG_DIR, spec["svg"])
        media = "ppt/media/" + spec["svg"]
        parts[media] = open(svg_src, "rb").read()

        rels_name = slide_name.replace("ppt/slides/", "ppt/slides/_rels/") + ".rels"
        rels = parts[rels_name].decode("utf-8")
        used = {int(m) for m in re.findall(r'Id="rId(\d+)"', rels)}
        rid = f"rId{max(used) + 1 if used else 1}"
        rels = rels.replace(
            "</Relationships>",
            f'<Relationship Id="{rid}" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
            f'Target="../media/{spec["svg"]}"/></Relationships>',
        )
        parts[rels_name] = rels.encode("utf-8")

        xml = parts[slide_name].decode("utf-8")
        # 図版は各スライドに1枚しか置いていないので、最初の <a:blip .../> を差し替える
        m = re.search(r'<a:blip r:embed="rId\d+"\s*/>', xml)
        if not m:
            raise RuntimeError(f"{slide_name}: <a:blip> が見つかりません")
        blip = m.group(0)[:-2] + ">"  # 自己終了タグを開きタグにする
        replacement = (
            blip
            + "<a:extLst>"
            + f'<a:ext uri="{SVG_EXT_URI}">'
            + f'<asvg:svgBlip xmlns:asvg="{SVG_NS}" r:embed="{rid}"/>'
            + "</a:ext></a:extLst></a:blip>"
        )
        xml = xml[: m.start()] + replacement + xml[m.end():]
        parts[slide_name] = xml.encode("utf-8")

    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, parts[n])
        for n in parts:
            if n not in names:
                zout.writestr(n, parts[n])
    os.replace(tmp, pptx_path)


def export_pngs(pptx_path):
    """.pptx → PDF → 1200×630 の PNG。スライドの並びと SLIDES の並びは同じ。"""
    run_soffice(["--convert-to", "pdf", "--outdir", WORK, pptx_path])
    pdf = os.path.join(WORK, os.path.splitext(os.path.basename(pptx_path))[0] + ".pdf")
    doc = fitz.open(pdf)
    if doc.page_count != len(SLIDES):
        raise RuntimeError(f"PDF のページ数が合いません: {doc.page_count}")
    os.makedirs(OUT_DIR, exist_ok=True)
    made = []
    for i, spec in enumerate(SLIDES):
        page = doc[i]
        zoom_x = 1200 / page.rect.width
        zoom_y = 630 / page.rect.height
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom_x, zoom_y), alpha=False)
        if (pix.width, pix.height) != (1200, 630):
            raise RuntimeError(f"寸法が違います: {pix.width}×{pix.height}")
        out = os.path.join(OUT_DIR, spec["png"])
        pix.save(out)
        made.append((out, os.path.getsize(out)))
    doc.close()
    return made


def main():
    shutil.rmtree(WORK, ignore_errors=True)
    os.makedirs(WORK, exist_ok=True)

    png_paths = {}
    for spec in SLIDES:
        src = os.path.join(SVG_DIR, spec["svg"])
        # PowerPoint 上の表示は 500×450 なので、代替 PNG は 2 倍で持たせる
        png_paths[spec["svg"]] = svg_to_png(src, WORK, 1000, 900)
        print(f"svg -> png: {spec['svg']}")

    build_pptx(png_paths)
    embed_svg(PPTX)
    print(f"pptx: {PPTX}")

    for out, size in export_pngs(PPTX):
        print(f"png : {out} ({size:,} bytes)")


if __name__ == "__main__":
    sys.exit(main())
